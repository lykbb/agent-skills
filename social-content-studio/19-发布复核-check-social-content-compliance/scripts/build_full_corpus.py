#!/usr/bin/env python3
"""Build the auditable corpus and deterministic lexicon from four source attachments."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paragraph_text(node: ET.Element) -> str:
    parts: list[str] = []
    for descendant in node.iter():
        if descendant.tag == f"{W_NS}t" and descendant.text:
            parts.append(descendant.text)
        elif descendant.tag == f"{W_NS}tab":
            parts.append("\t")
        elif descendant.tag in {f"{W_NS}br", f"{W_NS}cr"}:
            parts.append("\n")
    return "".join(parts).strip()


def extract_docx(path: Path, title: str) -> dict:
    records: list[dict] = []
    media: list[dict] = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        part_names = [
            name
            for name in names
            if re.fullmatch(
                r"word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml",
                name,
            )
        ]
        for part_name in sorted(part_names, key=lambda value: (value != "word/document.xml", value)):
            root = ET.fromstring(archive.read(part_name))
            part_ordinal = 0
            for paragraph in root.iter(f"{W_NS}p"):
                text = paragraph_text(paragraph)
                if not text:
                    continue
                part_ordinal += 1
                records.append(
                    {
                        "record_id": f"{path.stem}:{part_name}:{part_ordinal}",
                        "source_part": part_name,
                        "paragraph": part_ordinal,
                        "text": text,
                    }
                )
        for name in sorted(value for value in names if value.startswith("word/media/")):
            payload = archive.read(name)
            media.append(
                {
                    "name": name,
                    "size_bytes": len(payload),
                    "sha256": hashlib.sha256(payload).hexdigest(),
                }
            )
    return {
        "title": title,
        "local_input_name": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "record_count": len(records),
        "character_count": sum(len(item["text"]) for item in records),
        "records": records,
        "embedded_media": media,
    }


def extract_xlsx(path: Path, title: str) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets: list[dict] = []
    for sheet in workbook.worksheets:
        records: list[dict] = []
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                records.append(
                    {
                        "record_id": f"{path.stem}:{sheet.title}:{cell.coordinate}",
                        "sheet": sheet.title,
                        "coordinate": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                        "value": value.strip() if isinstance(value, str) else value,
                        "data_type": cell.data_type,
                        "number_format": cell.number_format,
                        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                        "comment": cell.comment.text if cell.comment else None,
                    }
                )
        sheets.append(
            {
                "title": sheet.title,
                "worksheet_dimension": {"max_row": sheet.max_row, "max_column": sheet.max_column},
                "last_content_row": max((item["row"] for item in records), default=0),
                "last_content_column": max((item["column"] for item in records), default=0),
                "merged_ranges": [str(value) for value in sheet.merged_cells.ranges],
                "record_count": len(records),
                "records": records,
            }
        )
    return {
        "title": title,
        "local_input_name": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "sheet_count": len(sheets),
        "record_count": sum(item["record_count"] for item in sheets),
        "sheets": sheets,
    }


def extract_pdf(path: Path, title: str) -> dict:
    reader = PdfReader(path)
    pages: list[dict] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        annotations = page.get("/Annots") or []
        pages.append(
            {
                "record_id": f"{path.stem}:page:{page_number}",
                "page": page_number,
                "text": text,
                "character_count": len(text),
                "image_count": len(getattr(page, "images", [])),
                "annotation_count": len(annotations),
                "visual_note": "封底，无新增规则" if page_number == 22 else None,
            }
        )
    return {
        "title": title,
        "local_input_name": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "page_count": len(pages),
        "nonempty_text_pages": sum(bool(item["text"]) for item in pages),
        "character_count": sum(item["character_count"] for item in pages),
        "pages": pages,
    }


COMMON_CONTEXT_TERMS = {
    "安全", "保证", "有效", "解决", "改善", "健康", "专业", "官方", "热销", "推荐",
    "第一", "最高", "最低", "最好", "最佳", "首个", "首款", "首选", "国家级",
    "世界级", "国际级", "领先", "高端", "先进", "天然", "科学", "独特", "权威",
    "消除", "提高", "增强", "促进", "修复", "抗菌", "杀菌", "减肥", "美白",
}


def category_for_column(records: list[dict], column: int) -> str:
    headers = [
        str(item["value"]).strip()
        for item in records
        if item["column"] == column and item["row"] in {1, 2} and item["value"] is not None
    ]
    return " / ".join(headers) if headers else f"第{column}列"


def split_terms(raw: str) -> list[str]:
    if len(raw) > 100:
        return []
    without_notes = re.sub(r"[（(][^（）()]{0,40}[）)]", "", raw)
    pieces = re.split(r"[、，,；;\n]|\s{2,}|/", without_notes)
    candidates: list[str] = []
    for piece in [without_notes, *pieces]:
        value = piece.strip(" \t。.!！?？:：'\"“”‘’[]【】")
        if not value or len(value) > 36:
            continue
        if value not in candidates:
            candidates.append(value)
    return candidates


def severity_for_category(category: str) -> str:
    if any(word in category for word in ("医疗", "金融", "国家形象", "消毒", "医疗器械", "保健食品")):
        return "high"
    if any(word in category for word in ("绝对化", "承诺", "封建迷信", "化妆品", "教育培训")):
        return "medium"
    return "review"


def compile_lexicon(spreadsheet: dict) -> dict:
    entries: list[dict] = []
    source_records = [record for sheet in spreadsheet["sheets"] for record in sheet["records"]]
    for record in source_records:
        if record["row"] < 4:
            continue
        raw = str(record["value"]).strip()
        category = category_for_column(source_records, record["column"])
        candidates = split_terms(raw)
        searchable = [value for value in candidates if len(value) >= 2]
        match_mode = "context" if not searchable or any(value in COMMON_CONTEXT_TERMS for value in searchable) else "phrase"
        entries.append(
            {
                "id": f"xlsx-{record['sheet']}-{record['coordinate']}",
                "platforms": ["xiaohongshu", "douyin"],
                "source_ref": f"违禁词.xlsx!{record['sheet']}!{record['coordinate']}",
                "source_record_id": record["record_id"],
                "category": category,
                "raw": raw,
                "terms": searchable,
                "match_mode": match_mode,
                "severity": severity_for_category(category),
                "context_note": "需结合商品类别、广告语境、限定范围和证明材料判断，命中不等于自动违规。",
            }
        )
    return {
        "version": f"{date.today().isoformat()}-four-source-full-lexicon-v2",
        "source": "违禁词.xlsx",
        "source_record_count_rows_4_plus": len(entries),
        "searchable_entry_count": sum(bool(item["terms"]) for item in entries),
        "context_only_entry_count": sum(item["match_mode"] == "context" for item in entries),
        "entries": entries,
    }


def behavior_rules() -> list[dict]:
    rows = [
        (1, "暴力演绎售卖", ["辱骂", "殴打", "重摔", "怒砸", "剪砍", "黑社会", "借贷剧情"], [5]),
        (2, "卖惨演绎售卖", ["哭喊叫卖", "编造悲惨故事", "弱势群体卖惨", "工厂倒闭", "破产停工"], [6]),
        (3, "炒作演绎售卖", ["反常识剧情", "不正当关系炒作", "夸张塞现金", "浮夸虚假表演"], [6]),
        (4, "错误表达爱国行为", ["恶意政治言论", "贬低他国", "过激情绪", "煽动民族情绪"], [6]),
        (5, "危险行为博眼球", ["吞食有毒活物", "过量进食", "无资质打针吃药", "高危挑战"], [6]),
        (6, "踩一捧一售卖", ["恶意贬低竞品", "故意小码对比", "夸张试吃表情", "片面对比"], [7]),
        (7, "恶性降价吆喝", ["虚构原价", "连续砍价", "夸张折扣", "价格与价值严重不符"], [7]),
        (8, "招募主播或代理", ["招募代理", "招商加盟", "轻松赚钱", "加入团队", "招聘导流"], [7]),
        (9, "虚假宠粉", ["评论送礼", "观看时长送礼", "虚假抽奖", "涨粉噱头", "未兑现福利"], [7]),
        (10, "私下交易与多渠道导流", ["主页联系方式", "线下交易", "绕过平台", "私下销售", "外部支付"], [8]),
        (11, "医疗描述", ["普通食品消炎", "治疗疾病", "奇效", "调节人体机能", "疗效承诺"], [8]),
        (12, "挂机直播或录播", ["单一画面", "长时间无交流", "主播离镜", "循环录播", "无版权成片"], [8]),
        (13, "衣着暴露", ["低领", "过度暴露", "不适宜着装"], [8]),
        (14, "语言过激", ["侮辱", "谩骂", "挑衅用户", "低俗攻击", "负面煽动"], [9]),
        (15, "吸烟饮酒", ["直播吸烟", "直播饮酒", "烟雾", "打火机声", "酒瓶展示"], [9]),
        (16, "言语低俗", ["低俗话题", "擦边话术", "低俗改词", "性暗示"], [9]),
        (17, "封建迷信", ["算命", "辟邪", "化解小人", "镇宅消灾", "轮回"], [9]),
        (18, "导师专家卖课牟利", ["导师卖课", "上热门保证", "月入六位数", "电商赚钱课", "课程收益保证"], [10]),
        (19, "未成年人直播带货", ["未成年人单独直播", "未成年人讲解商品", "未成年人推广", "十周岁以下代言"], [10]),
        (20, "低俗意味宣传商品", ["过度暴露售卖", "低俗动作", "被子裸露", "性化展示"], [10]),
        (21, "视频与商品无相关性", ["内容商品不一致", "购物车商品不相关", "挂错商品卡"], [10]),
        (22, "纹身长时间暴露", ["长时间展示纹身", "突出纹身"], [10]),
        (23, "毒品武器与管制器具", ["毒品", "管制刀具", "枪支", "仿真枪", "玩具枪"], [11]),
        (24, "境外直播高风险", ["境外直播", "旅拍直播带货", "海外开播"], [11]),
        (25, "截屏或评论区随机抽奖", ["截屏抽奖", "评论区抽奖", "随机送礼", "平台外抽奖"], [11]),
        (26, "通用或万能链接", ["一条链接卖多种商品", "拍1号链接", "万能链接", "商品与链接不符"], [11]),
        (27, "私信大量群发商品链接", ["群发链接", "私信轰炸", "批量私信商品"], [11]),
        (28, "好评返现", ["好评返现", "五星好评返钱", "评价后返现", "晒图返现"], [11]),
        (29, "宣传其他直播平台或非本人实名开播", ["其他平台直播", "多平台开播", "借用实名", "非本人实名认证"], [11]),
        (30, "冒充平台媒体或未授权官方活动", ["冒充官方", "冒充媒体", "未经授权活动", "虚假平台名义"], [11]),
        (31, "金钱营销与规避表达复核", ["高频金钱收益话术", "用米代替钱", "谐音规避", "诈骗式交易暗示"], [11]),
        (32, "诱骗秒杀", ["低价未上架", "秒杀未兑现", "价格不一致", "不限量免费送", "赠品未协商", "秒杀信息不完整"], [12, 13]),
        (33, "诱导互动", ["扣666领福利", "点赞换优惠", "停留时长换发货", "无意义评论换资格", "憋单"], [14]),
        (34, "其他严禁直播情况", ["事故现场直播", "无民事行为能力者单独直播", "乞讨直播", "军人直播", "恶搞名胜古迹", "游戏代练交易"], [15]),
    ]
    return [
        {
            "number": number,
            "id": f"douyin-behavior-{number:02d}",
            "title": title,
            "platforms": ["douyin"],
            "channels": ["口播", "字幕", "贴纸", "弹幕", "画面", "动作", "商品卡", "直播流程"],
            "severity": "high" if number not in {22, 24, 27, 31} else "review",
            "signals": signals,
            "safer_direction": "删除或改掉违规机制本身，真实展示商品、活动条件和平台内流程；不得用谐音、符号或拆词规避检测。",
            "source_ref": [f"抖音全套违禁词与违规行为规避指南.pdf#page={page}" for page in pages],
        }
        for number, title, signals, pages in rows
    ]


def curated_rules() -> list[dict]:
    both = ["xiaohongshu", "douyin"]
    return [
        {
            "id": "absolute-or-unverifiable-claims",
            "platforms": both,
            "severity": "high",
            "category": "绝对化、排名或无法证明的宣传",
            "terms": ["全网第一", "全球第一", "行业第一", "销量第一", "销量冠军", "唯一选择", "独一无二", "绝无仅有", "史无前例", "最好用", "最高级", "最先进", "全网最低价", "100%有效", "百分之百有效", "永久有效", "绝对有效", "绝对安全", "零风险"],
            "patterns": [r"(?:全网|全球|全国|行业|史上).{0,6}(?:第一|最好|最佳|最高|最低|唯一)", r"100\s*[%％].{0,8}(?:有效|安全|正品|成功|治愈|退款)", r"(?:保证|确保|承诺).{0,10}(?:有效|成功|学会|通过|治愈|赚钱)"],
            "reason": "排名、绝对效果、成功率或零风险承诺通常需要明确范围、时间和可核验证据，否则容易误导。",
            "suggestion": "删除绝对承诺；如确有证据，写明统计范围、时间、样本、条件和出处。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=2"],
        },
        {
            "id": "medical-health-body-claims",
            "platforms": both,
            "severity": "high",
            "category": "医疗、健康、减肥或人体机能宣称",
            "terms": ["治疗", "治愈", "根治", "医治", "临床有效", "消炎", "抗炎", "杀菌", "灭菌", "抗菌", "排毒", "解毒", "防癌", "抗癌", "降血压", "降血糖", "减肥", "瘦身", "溶脂", "吸脂", "燃烧脂肪", "毛囊激活", "提高免疫力", "增强免疫力", "祛疤", "脱敏", "改善睡眠", "补肾壮阳"],
            "patterns": [r"(?:\d+|[一二三四五六七八九十百千万]+)\s*天.{0,6}(?:瘦|减)\s*(?:\d+|[一二三四五六七八九十百千万]+)\s*斤", r"(?:\d+|[一二三四五六七八九十百千万]+)\s*(?:天|次|周).{0,6}(?:见效|治愈|根治|消除|解决)", r"(?:预防|防治).{0,10}(?:疾病|癌|高血压|糖尿病)"],
            "reason": "普通商品和普通内容不应承诺疾病治疗、固定期限见效或人体机能改变；特殊产品也只能在获批范围内宣传。",
            "suggestion": "改为真实成分、使用方法或个人使用记录；先确认产品资质和获批表述范围。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=4"],
        },
        {
            "id": "authority-certification-endorsement",
            "platforms": both,
            "severity": "high",
            "category": "权威、认证、政府或专家背书",
            "terms": ["国家级", "国家机关推荐", "国家机关专供", "领导人推荐", "质量免检", "无需检测", "免抽检", "中国驰名商标", "专家推荐", "官方认证", "军医推荐", "解放军推荐"],
            "patterns": [r"(?:国家|政府|领导人|部委|军队|专家|院士).{0,8}(?:推荐|认证|专供|指定|背书)"],
            "reason": "政府、军队、专家、认证或免检背书若无真实授权和资质，可能构成严重误导。",
            "suggestion": "删除无法证明的背书；如确有资质，展示可核验的完整名称、范围和有效期。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=3"],
        },
        {
            "id": "cosmetic-beauty-overclaim",
            "platforms": both,
            "severity": "medium",
            "category": "化妆品、美容或前后效果夸大",
            "terms": ["一洗白", "一抹白", "彻底美白", "彻底祛斑", "逆龄", "永葆青春", "100%修复", "永不复发", "永不褪色", "丰乳", "丰胸", "祛痘", "祛闭口", "破坏黑色素细胞"],
            "patterns": [r"(?:一洗|一抹|一次).{0,4}(?:白|祛斑|祛痘|见效)", r"(?:使用前|before).{0,10}(?:使用后|after)"],
            "reason": "化妆品效果、医美暗示和对比图容易越过产品宣称范围或造成误导。",
            "suggestion": "说明真实使用周期和个人感受，避免医疗化结论、固定效果和夸张前后对比。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=4"],
        },
        {
            "id": "finance-income-course-guarantee",
            "platforms": both,
            "severity": "high",
            "category": "金融收益、赚钱或课程结果保证",
            "terms": ["稳赚", "保本稳赚", "无风险回报", "固定回报", "保证收益", "包回报", "轻松赚钱", "月入六位数", "零成本赚钱", "保过", "保证就业", "保证通过考试", "上热门保证"],
            "patterns": [r"(?:月入|年入|日赚|周赚)\s*[十百千万亿0-9一二三四五六七八九]+(?:位数)?", r"(?:保证|包你|确保).{0,8}(?:赚钱|收益|回本|就业|通过|上岸|上热门)"],
            "reason": "收益、就业、考试、流量或回本保证容易构成欺骗性营销，金融类还可能触及更高风险。",
            "suggestion": "删除保证性结果，清楚说明成本、条件、失败可能和个体差异。",
            "source_refs": ["违禁词.xlsx", "抖音全套违禁词与违规行为规避指南.pdf#page=10"],
        },
        {
            "id": "price-scarcity-time-limit",
            "platforms": both,
            "severity": "medium",
            "category": "价格、时限、库存或稀缺性",
            "terms": ["史上最低价", "全网最低价", "马上涨价", "最后一波", "仅此一次", "再不抢就没了", "万人疯抢", "抢疯了", "抢爆了", "亏本卖", "倒贴卖", "随时结束", "随时涨价"],
            "patterns": [r"(?:限时|秒杀|特惠|优惠).{0,8}(?:无截止|随时结束|最后|仅此一次)", r"(?:库存|名额).{0,6}(?:只剩|最后)\s*[0-9一二三四五六七八九十]+"],
            "reason": "价格、时限、库存和折扣必须真实、完整、可兑现，并能说明比较依据。",
            "suggestion": "写清开始和结束时间、库存、适用条件、原价依据及预售情况。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=2"],
        },
        {
            "id": "deceptive-benefit-giveaway-flash-sale",
            "platforms": both,
            "severity": "high",
            "category": "虚假福利、抽奖或诱骗秒杀",
            "terms": ["不限量免费送", "人手一份", "全民免单", "恭喜获奖", "领取奖品", "截屏抽奖", "评论区随机抽奖", "秒杀未上架", "免费送未兑现"],
            "patterns": [r"(?:免费送|赠送|抽奖|秒杀).{0,18}(?:不限量|人手|人人|全部|百分百)", r"(?:秒杀|改价).{0,12}(?:未上架|不上架|未兑现|不兑现)"],
            "reason": "无法兑现的福利、未上架秒杀、信息不完整或平台外抽奖会损害交易体验。",
            "suggestion": "只发布已确认且能兑现的活动，完整展示时间、数量、条件、商品信息和规则。",
            "source_refs": ["抖音全套违禁词与违规行为规避指南.pdf#page=12", "抖音全套违禁词与违规行为规避指南.pdf#page=13"],
        },
        {
            "id": "induced-engagement",
            "platforms": both,
            "severity": "medium",
            "category": "诱导互动或把互动绑定利益",
            "terms": ["评论区扣1", "扣1领取", "扣666", "评论暗号", "一键三连", "点赞后发货", "评论后发货", "关注后领取", "停留5分钟", "拍了打三遍"],
            "patterns": [r"(?:扣|打|评论|刷)\s*(?:1|666|已买|拍了|想要).{0,12}(?:领取|报名|发货|优惠|福利|库存|资格)", r"(?:点赞|关注|评论|停留).{0,12}(?:换|领|获得|优先|才能).{0,10}(?:优惠|福利|发货|资格|包邮)"],
            "reason": "把无意义互动作为优惠、发货、改价或资格的前提，属于明显风险。",
            "suggestion": "取消互动前置条件；正常互动不要与虚假利益、差别发货或购买资格绑定。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=14"],
        },
        {
            "id": "private-contact-off-platform-transaction",
            "platforms": both,
            "severity": "high",
            "category": "私下联系、站外导流或绕过平台交易",
            "terms": ["加微信", "加我微信", "加V", "加薇", "私信我发链接", "关注后私信", "主页有联系方式", "线下交易", "站外下单", "外部二维码", "微信二维码", "淘口令", "绕过平台"],
            "patterns": [r"私信.{0,8}(?:发|给|拿).{0,8}(?:链接|二维码|联系方式|折扣)", r"(?:加|\+).{0,3}(?:微信|vx|v信|薇信|微\s*信)", r"(?:主页|简介).{0,8}(?:联系方式|二维码|链接)"],
            "reason": "引导私下联系、外部支付或绕开平台交易会显著增加处罚和诈骗风险。",
            "suggestion": "仅使用平台允许的商品卡、店铺、搜索和公开披露流程。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=8"],
        },
        {
            "id": "other-platform-mention-xhs",
            "platforms": ["xiaohongshu"],
            "severity": "review",
            "category": "其他平台名称或变体",
            "terms": ["淘宝", "天猫", "京东", "抖音", "快手", "拼多多", "微信", "微博", "B站", "T宝", "狗东", "某音", "某手", "多多", "小破站"],
            "reason": "中性提及、比较、来源说明与交易导流的风险不同，需要结合完整上下文。",
            "suggestion": "保留必要的中性事实；删除联系方式、外部购买指令和绕平台交易暗示。",
            "source_refs": ["小红书违禁词合规手册 (2).docx"],
        },
        {
            "id": "commercial-cooperation-disclosure",
            "platforms": both,
            "severity": "review",
            "category": "商业合作、广告或利益关系标注",
            "terms": ["品牌合作", "广告", "赞助", "推广", "佣金", "蒲公英", "商务合作", "置换"],
            "reason": "商业内容是否已按平台流程报备、标注，并与实际利益关系一致，需要人工核验。",
            "suggestion": "核对平台当前报备入口和标注要求，保留合作记录，避免伪装成无利益关系的自用分享。",
            "source_refs": ["小红书违禁词合规手册 (2).docx"],
        },
        {
            "id": "ai-content-disclosure",
            "platforms": both,
            "severity": "review",
            "category": "AI 生成或辅助内容标注",
            "terms": ["AI生成", "AI辅助", "数字人", "虚拟人", "换脸", "合成配音"],
            "reason": "AI 图片、文字、配音、数字人或换脸是否需要标注，取决于内容形式和平台当前规则。",
            "suggestion": "核对平台当前 AI 标注规则；需要时在发布设置和内容显著位置完整标注。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "小红书官方流量推送机制.docx"],
        },
        {
            "id": "originality-copyright-watermark",
            "platforms": both,
            "severity": "review",
            "category": "原创、版权、搬运或其他平台水印",
            "terms": ["搬运", "转载", "未经授权", "盗用", "抄袭", "他平台水印", "录屏直播", "无版权音乐", "素材来源不明"],
            "reason": "他人素材、音乐、脚本、画面、水印和高度雷同结构可能涉及版权或平台原创性风险。",
            "suggestion": "确认授权和来源，去除非必要水印，保留创作与许可凭证，进行实质性原创表达。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=16"],
        },
        {
            "id": "dangerous-illegal-content",
            "platforms": both,
            "severity": "high",
            "category": "武器、毒品、赌博、暴力、自伤或危险模仿",
            "terms": ["仿真枪", "玩具枪", "管制刀具", "毒品", "吸毒", "赌博", "暴力威胁", "自杀教程", "自残", "高危挑战", "吞食活物", "无资质打针"],
            "patterns": [r"(?:怒砸|砸|重摔|剪|砍).{0,8}(?:商品|竞品|产品|物品)"],
            "reason": "危险物、违法内容、可模仿伤害行为或高危演示通常属于高风险。",
            "suggestion": "删除危险展示和操作细节；必要的教育报道应避免可模仿步骤并提供清晰安全语境。",
            "source_refs": ["抖音全套违禁词与违规行为规避指南.pdf#page=6", "抖音全套违禁词与违规行为规避指南.pdf#page=11"],
        },
        {
            "id": "minor-sexual-tobacco-alcohol",
            "platforms": both,
            "severity": "high",
            "category": "未成年人、色情低俗、裸露、烟酒",
            "terms": ["未成年人带货", "儿童代言", "低俗动作", "性暗示", "裸露", "擦边", "吸烟", "饮酒", "直播喝酒", "直播抽烟"],
            "patterns": [r"未成年人.{0,8}(?:直播|带货|讲解|推广|口播|代言)"],
            "reason": "未成年人商业行为、性化展示、过度裸露和直播烟酒内容需要严格审查画面与参与方式。",
            "suggestion": "移除不适宜画面和动作；未成年人不得独立讲解或推广商品，并核对适龄和代言限制。",
            "source_refs": ["抖音全套违禁词与违规行为规避指南.pdf#page=8", "抖音全套违禁词与违规行为规避指南.pdf#page=10"],
        },
        {
            "id": "discrimination-politics-rumor",
            "platforms": both,
            "severity": "high",
            "category": "歧视、政治敏感、仇恨、谣言或极端煽动",
            "terms": ["民族歧视", "种族歧视", "性别歧视", "煽动仇恨", "政治谣言", "恶意政治言论", "邪教", "恐怖主义", "侮辱国家"],
            "reason": "仇恨、歧视、政治极端、谣言和煽动性内容可能触及平台与法律高风险。",
            "suggestion": "删除煽动和未经证实内容；必要报道使用权威来源并保持中性、完整语境。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=18"],
        },
        {
            "id": "superstition-pseudoscience",
            "platforms": both,
            "severity": "medium",
            "category": "迷信或伪科学营销",
            "terms": ["招财进宝", "化解小人", "提升运气", "逢凶化吉", "旺财", "护身符", "开光", "算命", "算卦", "镇宅消灾", "风水宝地", "增强第六感"],
            "reason": "把运势、风水、宗教或伪科学当作商品功效和事实保证存在明显风险。",
            "suggestion": "删除功效承诺；文化介绍要明确不是效果保证，不诱导付费消灾或改运。",
            "source_refs": ["违禁词.xlsx", "小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=9"],
        },
        {
            "id": "misleading-comparison-product-card",
            "platforms": both,
            "severity": "medium",
            "category": "恶意比较、虚假实验或商品不一致",
            "terms": ["踩一捧一", "恶意对比", "虚假实验", "夸张对比", "商品卡不一致", "万能链接", "一条链接卖多个商品", "伪造检测报告", "PS前后对比"],
            "reason": "片面对比、虚假实验、伪造证明、商品卡不匹配或通用链接会造成明显误导。",
            "suggestion": "使用同条件、可复现的真实比较；商品、链接、价格、参数和画面必须一致。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=7"],
        },
        {
            "id": "livestream-process-risk",
            "platforms": ["douyin"],
            "severity": "review",
            "category": "直播流程、录播、境外开播或实名问题",
            "terms": ["挂机直播", "录播", "无人直播", "境外直播", "多平台开播", "非本人实名认证", "冒充官方", "主播长时间离开", "循环播放"],
            "reason": "此类风险主要取决于实时画面、账号实名、开播地点和直播流程，不能只看口播文字。",
            "suggestion": "逐段核对真人出镜、实时互动、实名、地点、授权、商品卡和直播活动规则。",
            "source_refs": ["抖音全套违禁词与违规行为规避指南.pdf#page=8", "抖音全套违禁词与违规行为规避指南.pdf#page=11"],
        },
        {
            "id": "prohibited-regulated-products",
            "platforms": both,
            "severity": "high",
            "category": "药品、保健食品、医疗器械、医疗服务等受监管品类",
            "terms": ["OTC药品", "处方药", "保健食品", "蓝帽子", "医疗器械", "医美服务", "美白针", "溶脂针", "热玛吉", "超声刀", "在线诊疗", "注射手术"],
            "reason": "受监管产品能否推广、由谁推广、如何宣传和展示，取决于平台当前禁限售及资质规则。",
            "suggestion": "先核对品类准入、账号资质和平台最新禁限售清单；不确定时不要发布或带货。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "违禁词.xlsx"],
        },
        {
            "id": "privacy-impersonation-fake-identity",
            "platforms": both,
            "severity": "high",
            "category": "隐私、冒充或虚假身份资质",
            "terms": ["冒充官方", "冒充媒体", "虚构学历", "虚构职业", "网图冒充真人", "曝光手机号", "身份证", "住址", "个人二维码"],
            "patterns": [r"1[3-9]\d{9}", r"(?:身份证|住址|家庭地址).{0,12}[0-9a-zA-Z一二三四五六七八九十]"],
            "reason": "虚假身份、冒充机构和泄露个人敏感信息可能造成严重平台、隐私和欺诈风险。",
            "suggestion": "删除个人敏感信息和虚假身份，取得明确授权后再使用他人资料或机构名称。",
            "source_refs": ["小红书违禁词合规手册 (2).docx", "抖音全套违禁词与违规行为规避指南.pdf#page=11"],
        },
    ]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--traffic-docx", type=Path, required=True)
    parser.add_argument("--compliance-docx", type=Path, required=True)
    parser.add_argument("--terms-xlsx", type=Path, required=True)
    parser.add_argument("--douyin-pdf", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def write_json(path: Path, payload: dict | list) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    corpus = {
        "version": f"{date.today().isoformat()}-four-source-corpus-v2",
        "scope": "Complete text/cell/page extraction of the four user-supplied Feishu attachments.",
        "sources": {
            "xiaohongshu_traffic": extract_docx(args.traffic_docx, "小红书官方流量推送机制.docx"),
            "xiaohongshu_compliance": extract_docx(args.compliance_docx, "小红书违禁词合规手册 (2).docx"),
            "prohibited_terms": extract_xlsx(args.terms_xlsx, "违禁词.xlsx"),
            "douyin_guide": extract_pdf(args.douyin_pdf, "抖音全套违禁词与违规行为规避指南.pdf"),
        },
    }
    lexicon = compile_lexicon(corpus["sources"]["prohibited_terms"])
    behaviors = {
        "version": f"{date.today().isoformat()}-douyin-34-behaviors-v2",
        "source_boundary": "Secondary training material; use as a pre-publication risk checklist, not as proof of current official policy.",
        "rule_count": 34,
        "rules": behavior_rules(),
    }
    source_counts = {
        "xiaohongshu_traffic_records": corpus["sources"]["xiaohongshu_traffic"]["record_count"],
        "xiaohongshu_compliance_records": corpus["sources"]["xiaohongshu_compliance"]["record_count"],
        "xlsx_nonempty_cells": corpus["sources"]["prohibited_terms"]["record_count"],
        "xlsx_content_rows": corpus["sources"]["prohibited_terms"]["sheets"][0]["last_content_row"],
        "xlsx_compiled_records_rows_4_plus": lexicon["source_record_count_rows_4_plus"],
        "pdf_pages": corpus["sources"]["douyin_guide"]["page_count"],
        "pdf_nonempty_text_pages": corpus["sources"]["douyin_guide"]["nonempty_text_pages"],
        "douyin_behavior_rules": len(behaviors["rules"]),
    }
    coverage = {
        "version": f"{date.today().isoformat()}-coverage-v2",
        "source_counts": source_counts,
        "source_sha256": {
            key: value["sha256"]
            for key, value in corpus["sources"].items()
        },
        "checks": {
            "all_xlsx_cells_rows_4_plus_have_lexicon_records": lexicon["source_record_count_rows_4_plus"]
            == sum(
                record["row"] >= 4
                for sheet in corpus["sources"]["prohibited_terms"]["sheets"]
                for record in sheet["records"]
            ),
            "all_34_douyin_behavior_numbers_present": [item["number"] for item in behaviors["rules"]]
            == list(range(1, 35)),
            "pdf_page_22_visually_reviewed_as_back_cover": corpus["sources"]["douyin_guide"]["pages"][21]["visual_note"]
            == "封底，无新增规则",
        },
        "scope_statement": "Complete relative to the four supplied files; platform rules can change and official live review still controls.",
    }
    curated = {
        "version": f"{date.today().isoformat()}-four-source-curated-rules-v2",
        "rules": curated_rules(),
    }
    write_json(args.output_dir / "source-corpus.json", corpus)
    write_json(args.output_dir / "full-terms.json", lexicon)
    write_json(args.output_dir / "behavior-rules.json", behaviors)
    write_json(args.output_dir / "coverage.json", coverage)
    write_json(args.output_dir / "rules.json", curated)
    print(json.dumps({"outputs": 5, **source_counts, "curated_rules": len(curated["rules"]), "checks": coverage["checks"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
